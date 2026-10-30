from __future__ import annotations

import hashlib
import hmac
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import tempfile
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from typing import Annotated

import imageio_ffmpeg
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from dotenv import load_dotenv
from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from openai import OpenAI
from PIL import Image


ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
DATA_ROOT = Path(os.getenv("FIO_DATA_DIR", ROOT / "data")).resolve()
DB_PATH = DATA_ROOT / "fio.sqlite3"
UPLOAD_ROOT = DATA_ROOT / "uploads"
OUTPUT_ROOT = DATA_ROOT / "documents"
MAX_UPLOAD_BYTES = int(os.getenv("FIO_MAX_UPLOAD_MB", "500")) * 1024 * 1024
MAX_UNPACKED_BYTES = int(os.getenv("FIO_MAX_UNPACKED_MB", "2048")) * 1024 * 1024
COOKIE_NAME = "fio_session"
COOKIE_SECURE = os.getenv("FIO_COOKIE_SECURE", "false").lower() == "true"
COOKIE_SAMESITE = os.getenv("FIO_COOKIE_SAMESITE", "lax").lower()
ALLOWED_ORIGINS = [origin.strip() for origin in os.getenv(
    "FIO_ALLOWED_ORIGINS",
    "http://localhost:5173,http://127.0.0.1:5173,https://localhost,capacitor://localhost",
).split(",") if origin.strip()]
TRANSCRIPTION_MODEL = os.getenv("FIO_TRANSCRIPTION_MODEL", "gpt-transcribe")
SESSION_DAYS = int(os.getenv("FIO_SESSION_DAYS", "7"))
MONTHS = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
HEADER_RE = re.compile(r"^(\d{2}/\d{2}/\d{4}) (\d{2}:\d{2}) - (.*)$")
IMAGE_RE = re.compile(r"((?:IMG|STK)-[^\s<>]+\.(?:jpg|jpeg|png|webp))", re.IGNORECASE)
AUDIO_RE = re.compile(r"\b([A-Za-z0-9_][A-Za-z0-9_.-]*\.(?:opus|ogg|m4a|mp3|aac|amr|wav))\b", re.IGNORECASE)

DATA_ROOT.mkdir(parents=True, exist_ok=True)
UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="DocumentAI API", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type"],
)


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def db() -> sqlite3.Connection:
    connection = sqlite3.connect(DB_PATH, timeout=15)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def init_db() -> None:
    with db() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                password_salt TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token_hash TEXT PRIMARY KEY,
                user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                expires_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS jobs (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                original_name TEXT NOT NULL,
                status TEXT NOT NULL,
                progress TEXT NOT NULL,
                error TEXT,
                output_path TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS jobs_user_created ON jobs(user_id, created_at DESC);
            """
        )


init_db()


def hash_password(password: str, salt: bytes | None = None) -> tuple[str, str]:
    actual_salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), actual_salt, 390_000)
    return digest.hex(), actual_salt.hex()


def verify_password(password: str, salt_hex: str, expected_hex: str) -> bool:
    actual, _ = hash_password(password, bytes.fromhex(salt_hex))
    return hmac.compare_digest(actual, expected_hex)


def issue_session(response: Response, user_id: str) -> None:
    token = secrets.token_urlsafe(48)
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    expires = utc_now() + timedelta(days=SESSION_DAYS)
    with db() as con:
        con.execute(
            "INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)",
            (token_hash, user_id, expires.isoformat()),
        )
    response.set_cookie(
        COOKIE_NAME,
        token,
        httponly=True,
        secure=COOKIE_SECURE,
        samesite=COOKIE_SAMESITE,
        max_age=SESSION_DAYS * 24 * 60 * 60,
        path="/",
    )


def current_user(request: Request) -> dict:
    token = request.cookies.get(COOKIE_NAME)
    if not token:
        raise HTTPException(401, "Entre ou crie uma conta para continuar.")
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    with db() as con:
        row = con.execute(
            "SELECT u.id,u.name,u.email,s.expires_at FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=?",
            (token_hash,),
        ).fetchone()
        if not row:
            raise HTTPException(401, "Sessão expirada. Entre novamente.")
        if datetime.fromisoformat(row["expires_at"]) <= utc_now():
            con.execute("DELETE FROM sessions WHERE token_hash=?", (token_hash,))
            raise HTTPException(401, "Sessão expirada. Entre novamente.")
    return dict(row)


CurrentUser = Annotated[dict, Depends(current_user)]


@app.get("/api/health")
def health() -> dict:
    return {"status": "ok", "transcription_configured": bool(os.getenv("OPENAI_API_KEY"))}


@app.post("/api/auth/register")
def register(
    response: Response,
    name: Annotated[str, Form(min_length=1, max_length=100)],
    email: Annotated[str, Form(min_length=3, max_length=254)],
    password: Annotated[str, Form(min_length=8, max_length=200)],
) -> dict:
    clean_email = email.strip().lower()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", clean_email):
        raise HTTPException(422, "Informe um e-mail válido.")
    password_hash, password_salt = hash_password(password)
    user_id = str(uuid.uuid4())
    try:
        with db() as con:
            con.execute(
                "INSERT INTO users(id,name,email,password_hash,password_salt,created_at) VALUES(?,?,?,?,?,?)",
                (user_id, name.strip(), clean_email, password_hash, password_salt, utc_now().isoformat()),
            )
    except sqlite3.IntegrityError:
        raise HTTPException(409, "Já existe uma conta com esse e-mail.")
    issue_session(response, user_id)
    return {"id": user_id, "name": name.strip(), "email": clean_email}


@app.post("/api/auth/login")
def login(
    response: Response,
    email: Annotated[str, Form()],
    password: Annotated[str, Form()],
) -> dict:
    clean_email = email.strip().lower()
    with db() as con:
        user = con.execute("SELECT * FROM users WHERE email=?", (clean_email,)).fetchone()
    if not user or not verify_password(password, user["password_salt"], user["password_hash"]):
        raise HTTPException(401, "E-mail ou senha inválidos.")
    issue_session(response, user["id"])
    return {"id": user["id"], "name": user["name"], "email": user["email"]}


@app.post("/api/auth/logout")
def logout(request: Request, response: Response) -> dict:
    token = request.cookies.get(COOKIE_NAME)
    if token:
        with db() as con:
            con.execute("DELETE FROM sessions WHERE token_hash=?", (hashlib.sha256(token.encode()).hexdigest(),))
    response.delete_cookie(COOKIE_NAME, path="/")
    return {"ok": True}


@app.get("/api/me")
def me(user: CurrentUser) -> dict:
    return {"id": user["id"], "name": user["name"], "email": user["email"]}


def update_job(job_id: str, status: str, progress: str, error: str | None = None, output_path: str | None = None) -> None:
    with db() as con:
        con.execute(
            "UPDATE jobs SET status=?,progress=?,error=?,output_path=COALESCE(?,output_path),updated_at=? WHERE id=?",
            (status, progress, error, output_path, utc_now().isoformat(), job_id),
        )


async def save_upload(upload: UploadFile, destination: Path) -> int:
    total = 0
    with destination.open("wb") as output:
        while chunk := await upload.read(1024 * 1024):
            total += len(chunk)
            if total > MAX_UPLOAD_BYTES:
                destination.unlink(missing_ok=True)
                raise HTTPException(413, f"O ZIP ultrapassa o limite de {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
            output.write(chunk)
    return total


def safe_extract(archive: zipfile.ZipFile, destination: Path) -> None:
    root = destination.resolve()
    infos = archive.infolist()
    total = sum(info.file_size for info in infos)
    if total > MAX_UNPACKED_BYTES:
        raise ValueError("O conteúdo descompactado excede o limite permitido.")
    for info in infos:
        if info.is_dir():
            continue
        normalized = info.filename.replace("\\", "/")
        relative = PurePosixPath(normalized)
        if relative.is_absolute() or any(part in ("", ".", "..") for part in relative.parts):
            raise ValueError("O ZIP contém um caminho de arquivo inválido.")
        target = (root / Path(*relative.parts)).resolve()
        if root not in target.parents:
            raise ValueError("O ZIP contém um caminho fora da pasta do trabalho.")
        target.parent.mkdir(parents=True, exist_ok=True)
        written = 0
        with archive.open(info) as source, target.open("wb") as output:
            while chunk := source.read(1024 * 1024):
                written += len(chunk)
                if written > info.file_size:
                    raise ValueError("O ZIP contém um item com tamanho inválido.")
                output.write(chunk)


def parse_chat(chat_file: Path) -> list[dict[str, str]]:
    raw = chat_file.read_text(encoding="utf-8-sig", errors="replace")
    messages: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for line in raw.splitlines():
        match = HEADER_RE.match(line)
        if match:
            if current is not None:
                messages.append(current)
            date, time, rest = match.groups()
            if ": " in rest:
                sender, body = rest.split(": ", 1)
            else:
                sender, body = "Sistema", rest
            current = {"date": date, "time": time, "sender": sender, "body": body}
        elif current is not None:
            current["body"] += "\n" + line
    if current is not None:
        messages.append(current)
    if not messages:
        raise ValueError("Não encontrei mensagens no TXT exportado do WhatsApp.")
    return messages


def media_index(folder: Path) -> dict[str, Path]:
    return {file.name.casefold(): file for file in folder.rglob("*") if file.is_file()}


def transcribe_audio(audio: Path, scratch: Path) -> str:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("A transcrição automática foi autorizada, mas o provedor ainda não está configurado.")
    converted = scratch / (audio.stem + ".mp3")
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    result = subprocess.run(
        [ffmpeg, "-y", "-v", "error", "-i", str(audio), "-vn", "-ac", "1", "-ar", "16000", "-b:a", "40k", str(converted)],
        capture_output=True,
        timeout=180,
    )
    if result.returncode != 0 or not converted.exists():
        raise RuntimeError(f"Não consegui preparar o áudio {audio.name} para transcrição.")
    if converted.stat().st_size > 24 * 1024 * 1024:
        raise RuntimeError(f"O áudio {audio.name} é longo demais para a transcrição automática atual.")
    client = OpenAI(api_key=api_key)
    with converted.open("rb") as audio_file:
        response = client.audio.transcriptions.create(
            model=TRANSCRIPTION_MODEL,
            file=audio_file,
            language="pt",
        )
    return response.text.strip()


def set_cell_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def page_field(run, name: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def build_docx(messages: list[dict[str, str]], media: dict[str, Path], transcripts: dict[str, str], work: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.68)
    section.bottom_margin = Inches(.68)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(35, 40, 38)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 21), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(25, 35, 45)
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 2"].paragraph_format.space_before = Pt(11)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(5)

    document.add_paragraph("Exportação da conversa do WhatsApp com anexos", style="Title")
    p = document.add_paragraph("Mensagens em ordem cronológica, com imagens e transcrições automáticas autorizadas.")
    p.paragraph_format.space_after = Pt(8)
    intro = document.add_paragraph(
        "As mensagens e os anexos são apresentados conforme a exportação. A transcrição automática pode conter erros; confira o texto com o áudio original. Este documento organiza o material e não certifica sua autenticidade."
    )
    intro.paragraph_format.space_after = Pt(14)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Conversa exportada do WhatsApp  •  Página ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(110, 120, 130)
    page_field(footer.add_run(), "PAGE")
    footer.add_run(" de ")
    page_field(footer.add_run(), "NUMPAGES")

    last_date = None
    for message in messages:
        if message["date"] != last_date:
            day, month, year = map(int, message["date"].split("/"))
            heading = document.add_heading(f"{day} de {MONTHS[month - 1]} de {year}", level=2)
            p_pr = heading._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "3")
            bottom.set(qn("w:color"), "C9CDD2")
            borders.append(bottom)
            p_pr.append(borders)
            last_date = message["date"]

        body = document.add_paragraph()
        body.paragraph_format.space_after = Pt(5)
        body.paragraph_format.keep_together = True
        time_run = body.add_run(f"[{message['time']}] ")
        time_run.font.color.rgb = RGBColor(95, 100, 105)
        sender_run = body.add_run(message["sender"] + ": ")
        sender_run.bold = True
        for index, line in enumerate(message["body"].split("\n")):
            if index:
                body.add_run().add_break()
            body.add_run(line)

        image_names = [match.group(1) for match in IMAGE_RE.finditer(message["body"])]
        for name in image_names:
            image = media.get(name.casefold())
            if not image:
                continue
            image_for_doc = image
            with Image.open(image) as original:
                width, height = original.size
                if image.suffix.casefold() == ".webp":
                    image_for_doc = work / (image.stem + "_word.png")
                    original.convert("RGBA").save(image_for_doc, format="PNG")
            scale = min(5.2 / width, 7.2 / height)
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_before = Pt(2)
            picture.paragraph_format.space_after = Pt(1)
            picture.add_run().add_picture(str(image_for_doc), width=Inches(width * scale), height=Inches(height * scale))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(7)
            cap = caption.add_run("Imagem anexada — " + name)
            cap.italic = True
            cap.font.size = Pt(8)
            cap.font.color.rgb = RGBColor(100, 110, 120)

        audio_names = [match.group(1).strip() for match in AUDIO_RE.finditer(message["body"])]
        for name in dict.fromkeys(audio_names):
            label = document.add_paragraph()
            label.paragraph_format.left_indent = Inches(.24)
            label.paragraph_format.keep_with_next = True
            label.paragraph_format.space_after = Pt(2)
            label_run = label.add_run("[Transcrição automática — " + name + "]")
            label_run.bold = True
            label_run.font.size = Pt(8.5)
            label_run.font.color.rgb = RGBColor(70, 80, 90)
            transcript = transcripts.get(name.casefold())
            transcript = transcript if transcript else "Transcrição não solicitada para este áudio."
            for segment in transcript.splitlines() or [transcript]:
                if not segment.strip():
                    continue
                text_paragraph = document.add_paragraph(segment)
                text_paragraph.paragraph_format.left_indent = Inches(.24)
                text_paragraph.paragraph_format.right_indent = Inches(.12)
                text_paragraph.paragraph_format.space_after = Pt(3)

    document.core_properties.title = "Exportação da conversa do WhatsApp com anexos"
    document.core_properties.subject = "Histórico cronológico com imagens e transcrições automáticas"
    document.core_properties.author = "DocumentAI"
    document.save(output)


def process_job(job_id: str, user_id: str, archive_path: Path, transcription_consent: bool) -> None:
    work = Path(tempfile.mkdtemp(prefix=f"fio-{job_id}-", dir=UPLOAD_ROOT))
    try:
        update_job(job_id, "processing", "Conferindo o arquivo ZIP")
        with zipfile.ZipFile(archive_path) as archive:
            safe_extract(archive, work)
        chat_files = [file for file in work.rglob("*.txt") if file.is_file()]
        if not chat_files:
            raise ValueError("O ZIP não contém o TXT exportado do WhatsApp.")
        chat_file = max(chat_files, key=lambda file: file.stat().st_size)
        messages = parse_chat(chat_file)
        media = media_index(work)
        audio_names: list[str] = []
        for message in messages:
            audio_names.extend(match.group(1).strip() for match in AUDIO_RE.finditer(message["body"]))
        audio_names = list(dict.fromkeys(audio_names))
        transcripts: dict[str, str] = {}
        if transcription_consent:
            if not os.getenv("OPENAI_API_KEY"):
                raise RuntimeError("A transcrição ainda não está configurada. Nenhum áudio foi enviado.")
            for index, name in enumerate(audio_names, start=1):
                audio = media.get(name.casefold())
                if not audio:
                    continue
                update_job(job_id, "processing", f"Transcrevendo áudio {index} de {len(audio_names)}")
                transcripts[name.casefold()] = transcribe_audio(audio, work).strip()
        update_job(job_id, "processing", "Montando o documento Word")
        user_output = OUTPUT_ROOT / user_id
        user_output.mkdir(parents=True, exist_ok=True)
        output = user_output / f"{job_id}.docx"
        build_docx(messages, media, transcripts, work, output)
        update_job(job_id, "ready", "Documento pronto para baixar", output_path=str(output))
    except Exception as error:
        safe_message = str(error)[:500] or "Não foi possível processar esta conversa."
        update_job(job_id, "failed", "Não foi possível gerar o documento", error=safe_message)
    finally:
        archive_path.unlink(missing_ok=True)
        shutil.rmtree(work, ignore_errors=True)


@app.post("/api/jobs", status_code=202)
async def create_job(
    background_tasks: BackgroundTasks,
    user: CurrentUser,
    archive: Annotated[UploadFile, File()],
    transcription_consent: Annotated[bool, Form()] = False,
    delete_source_after_export: Annotated[bool, Form()] = True,
) -> dict:
    if not archive.filename or not archive.filename.lower().endswith(".zip"):
        raise HTTPException(415, "Envie um arquivo ZIP exportado pelo WhatsApp.")
    job_id = str(uuid.uuid4())
    job_root = UPLOAD_ROOT / user["id"]
    job_root.mkdir(parents=True, exist_ok=True)
    archive_path = job_root / f"{job_id}.zip"
    await save_upload(archive, archive_path)
    if not zipfile.is_zipfile(archive_path):
        archive_path.unlink(missing_ok=True)
        raise HTTPException(422, "O arquivo escolhido não é um ZIP válido.")
    now = utc_now().isoformat()
    with db() as con:
        con.execute(
            "INSERT INTO jobs(id,user_id,original_name,status,progress,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
            (job_id, user["id"], Path(archive.filename).name, "queued", "Na fila", now, now),
        )
    # The policy is always ephemeral for source uploads, even if a modified client omits the flag.
    del_source = delete_source_after_export
    if del_source:
        background_tasks.add_task(process_job, job_id, user["id"], archive_path, transcription_consent)
    else:
        archive_path.unlink(missing_ok=True)
        update_job(job_id, "failed", "A política de retenção exige apagar o ZIP de origem.", error="Reenvie com a exclusão dos arquivos originais ativada.")
    return {"id": job_id, "status": "queued", "progress": "Na fila"}


def job_view(row: sqlite3.Row) -> dict:
    return {
        "id": row["id"],
        "name": row["original_name"],
        "status": row["status"],
        "progress": row["progress"],
        "error": row["error"],
        "created_at": row["created_at"],
        "download_url": f"/api/jobs/{row['id']}/download" if row["status"] == "ready" else None,
    }


@app.get("/api/jobs")
def list_jobs(user: CurrentUser) -> list[dict]:
    with db() as con:
        rows = con.execute("SELECT * FROM jobs WHERE user_id=? ORDER BY created_at DESC", (user["id"],)).fetchall()
    return [job_view(row) for row in rows]


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str, user: CurrentUser) -> dict:
    with db() as con:
        row = con.execute("SELECT * FROM jobs WHERE id=? AND user_id=?", (job_id, user["id"])).fetchone()
    if not row:
        raise HTTPException(404, "Documento não encontrado.")
    return job_view(row)


@app.get("/api/jobs/{job_id}/download")
def download_job(job_id: str, user: CurrentUser) -> FileResponse:
    with db() as con:
        row = con.execute("SELECT * FROM jobs WHERE id=? AND user_id=?", (job_id, user["id"])).fetchone()
    if not row or row["status"] != "ready" or not row["output_path"]:
        raise HTTPException(404, "O documento ainda não está pronto para baixar.")
    output = Path(row["output_path"]).resolve()
    if not output.is_file() or OUTPUT_ROOT.resolve() not in output.parents:
        raise HTTPException(404, "O arquivo não está disponível.")
    name = Path(row["original_name"]).stem + "_organizado.docx"
    return FileResponse(output, filename=name, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.delete("/api/jobs/{job_id}")
def delete_job(job_id: str, user: CurrentUser) -> dict:
    with db() as con:
        row = con.execute("SELECT output_path FROM jobs WHERE id=? AND user_id=?", (job_id, user["id"])).fetchone()
        if not row:
            raise HTTPException(404, "Documento não encontrado.")
        if row["output_path"]:
            output = Path(row["output_path"]).resolve()
            if OUTPUT_ROOT.resolve() in output.parents:
                output.unlink(missing_ok=True)
        con.execute("DELETE FROM jobs WHERE id=? AND user_id=?", (job_id, user["id"]))
    return {"ok": True}
